"""DOCX text extraction using python-docx.

python-docx has no reliable page-layout information, so extracted sections
carry no page numbers. Paragraph and table text are both included because table
content is often where key facts live.

Visual detection looks for paragraphs carrying ``w:drawing`` elements (inline
and anchored images are both wrapped in ``w:drawing``). The containing
paragraph's text becomes the caption, with the next non-empty paragraph added
as surrounding context. As with every extractor this is best-effort: a hostile
document can only suppress visual metadata, never break text extraction.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from rag.errors import CorruptDocumentError, MissingFileError
from rag.extractors.base import Extractor
from rag.extractors.visuals import guess_visual_type
from rag.models import (
    VISUAL_TYPE_IMAGE,
    ExtractedDocument,
    ExtractedSection,
    VisualElement,
)
from rag.normalize import normalize_text


class DocxExtractor(Extractor):
    file_type = ".docx"

    def extract(self, source: str | Path) -> ExtractedDocument:
        path = self.resolve_path(source)
        try:
            document = Document(str(path))
        except (PackageNotFoundError, zipfile.BadZipFile) as exc:
            raise CorruptDocumentError(
                "The DOCX document is corrupt or not a valid Word file"
            ) from exc
        except OSError as exc:
            raise MissingFileError("The source document could not be read") from exc

        parts: list[str] = []
        visuals: list[VisualElement] = []

        paragraphs = list(document.paragraphs)
        for index, paragraph in enumerate(paragraphs):
            if paragraph.text and paragraph.text.strip():
                parts.append(paragraph.text.strip())

            drawing = self._paragraph_has_drawing(paragraph)
            if not drawing:
                continue
            own = (paragraph.text or "").strip()
            following = ""
            for candidate in paragraphs[index + 1 :]:
                if candidate.text and candidate.text.strip():
                    following = candidate.text.strip()
                    break
            caption = normalize_text(own or following)
            nearby_text = normalize_text(
                "\n".join(part for part in (own, following) if part)
            )
            visuals.append(
                VisualElement(
                    visual_id=f"inline{index}",
                    source_location=None,
                    visual_type=guess_visual_type(caption, fallback=VISUAL_TYPE_IMAGE),
                    caption=caption,
                    nearby_text=nearby_text,
                )
            )

        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))

        text = normalize_text("\n".join(parts))
        return ExtractedDocument(
            file_type=self.file_type,
            text=text,
            sections=(ExtractedSection(text=text),),
            visuals=tuple(visuals),
        )

    @staticmethod
    def _paragraph_has_drawing(paragraph) -> bool:
        try:
            return bool(paragraph._p.xpath(".//w:drawing"))
        except Exception:
            return False