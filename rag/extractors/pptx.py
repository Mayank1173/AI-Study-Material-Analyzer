"""PPTX text extraction using python-pptx, slide by slide.

Every slide becomes one addressable section carrying its slide number so chunk
metadata can point back to the exact slide the text came from. Speaker notes
are included when present because they often carry the slide's talking points.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.exc import PackageNotFoundError

from rag.errors import CorruptDocumentError, MissingFileError
from rag.extractors.base import Extractor
from rag.models import ExtractedDocument, ExtractedSection
from rag.normalize import normalize_text


class PptxExtractor(Extractor):
    file_type = ".pptx"

    def extract(self, source: str | Path) -> ExtractedDocument:
        path = self.resolve_path(source)
        try:
            presentation = Presentation(str(path))
        except (PackageNotFoundError, zipfile.BadZipFile) as exc:
            raise CorruptDocumentError(
                "The PPTX document is corrupt or not a valid PowerPoint file"
            ) from exc
        except OSError as exc:
            raise MissingFileError("The source document could not be read") from exc

        sections: list[ExtractedSection] = []
        for index, slide in enumerate(presentation.slides, start=1):
            lines = self._slide_text(slide)
            text = normalize_text("\n".join(lines))
            sections.append(
                ExtractedSection(
                    text=text,
                    slide=index,
                    source_location=f"slide {index}",
                )
            )

        return ExtractedDocument(
            file_type=self.file_type,
            text="\n\n".join(
                section.text for section in sections if section.text
            ),
            sections=tuple(sections),
        )

    @staticmethod
    def _slide_text(slide) -> list[str]:
        lines: list[str] = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                content = "".join(run.text for run in paragraph.runs).strip()
                if content:
                    lines.append(content)

        try:
            notes = slide.notes_slide.notes_text_frame.text
        except (KeyError, ValueError, AttributeError):
            notes = ""
        if notes and notes.strip():
            lines.append(f"[notes] {notes.strip()}")
        return lines