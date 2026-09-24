"""Phase 5 tests: diagram & visual content intelligence.

Covers visual element extraction from PDF/PPTX/DOCX, visual persistence and
keyword retrieval in the vector store, visual-query understanding, grounded
visual prompts with continued [Source N] numbering, security scoping, and the
best-effort failure guarantees that keep visual analysis from ever breaking
text extraction or answer generation.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from rag.answer import answer_question, sources_from_visuals
from rag.context_builder import (
    CONTEXT_DELIMITER_END,
    CONTEXT_DELIMITER_START,
    NO_CONTEXT_SYSTEM_PROMPT,
    SYSTEM_PROMPT,
    VISUAL_GROUNDING_NOTICE,
    build_grounded_prompt,
)
from rag.embeddings import DeterministicEmbedder
from rag.evidence import collect_visual_evidence
from rag.intent import analyze_query
from rag.knowledge_base import KnowledgeBase
from rag.llm.mock_provider import MockProvider
from rag.models import (
    VISUAL_TYPE_CHART,
    VISUAL_TYPE_FIGURE,
    VISUAL_TYPE_IMAGE,
    ChunkMetadata,
    SearchResult,
    SourceRef,
    VisualElement,
)
from rag.pipeline import process_document
from rag.tests.sample_data import PAGE_ONE_BODY, PAGE_ONE_HEADING
from rag.vectorstore import SqliteVectorStore
from rag.vectorstore.base import SearchFilter


def _write_png(path: Path) -> None:
    image = Image.new("RGB", (24, 24), (180, 80, 80))
    image.save(str(path), format="PNG")


@pytest.fixture()
def figure_pdf_path(tmp_path: Path) -> Path:
    _write_png(tmp_path / "fig1.png")
    path = tmp_path / "figure.pdf"
    pdf = canvas.Canvas(str(path), pagesize=letter)
    pdf.drawString(72, 720, PAGE_ONE_HEADING)
    pdf.drawString(72, 696, PAGE_ONE_BODY)
    pdf.showPage()
    pdf.drawString(72, 720, "Figure 1: Deadlock Detection Algorithm")
    pdf.drawString(72, 696, "The algorithm builds a wait-for graph and scans its edges.")
    pdf.drawImage(str(tmp_path / "fig1.png"), 100, 300, 60, 60)
    pdf.showPage()
    pdf.save()
    return path


@pytest.fixture()
def picture_pptx_path(tmp_path: Path) -> Path:
    _write_png(tmp_path / "pic1.png")
    path = tmp_path / "figure_slides.pptx"
    presentation = Presentation()
    blank = presentation.slide_layouts[6]

    slide_one = presentation.slides.add_slide(blank)
    title = slide_one.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(8), Inches(0.8))
    title.text = "Deadlock Detection"
    slide_one.shapes.add_picture(
        str(tmp_path / "pic1.png"), Inches(1), Inches(1.5), width=Inches(2)
    )

    slide_two = presentation.slides.add_slide(blank)
    title_two = slide_two.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(8), Inches(0.8))
    title_two.text = "Banker's Algorithm"
    chart_data = CategoryChartData()
    chart_data.categories = ["A", "B"]
    chart_data.add_series("Allocation", (1, 2))
    slide_two.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(1),
        Inches(1.5),
        Inches(3),
        Inches(2),
        chart_data,
    )

    presentation.save(str(path))
    return path


@pytest.fixture()
def inline_image_docx_path(tmp_path: Path) -> Path:
    _write_png(tmp_path / "inline.png")
    path = tmp_path / "figure_notes.docx"
    document = Document()
    document.add_paragraph(PAGE_ONE_HEADING)
    document.add_picture(str(tmp_path / "inline.png"))
    document.add_paragraph("Figure 2: Resource Allocation Graph explained.")
    document.save(str(path))
    return path


@pytest.fixture()
def embedder() -> DeterministicEmbedder:
    return DeterministicEmbedder(dimension=64)


@pytest.fixture()
def kb(embedder: DeterministicEmbedder, tmp_path: Path) -> KnowledgeBase:
    base = KnowledgeBase(embedder, SqliteVectorStore(str(tmp_path / "visual.db")))
    yield base
    base.close()


def _visual(
    visual_id: str = "m1:p2",
    material_id: str = "m1",
    course_id: str = "c1",
    page: int | None = 12,
    source_location: str | None = "page 12",
    visual_type: str = "diagram",
    caption: str = "Deadlock Detection Algorithm",
    nearby_text: str = "The wait-for graph holds one node per process.",
    **kwargs,
) -> VisualElement:
    return VisualElement(
        visual_id=visual_id,
        material_id=material_id,
        course_id=course_id,
        page=page,
        source_location=source_location,
        visual_type=visual_type,
        caption=caption,
        nearby_text=nearby_text,
        **kwargs,
    )


class _FakeKnowledgeBase:
    def __init__(
        self,
        results: list[SearchResult] | None = None,
        visuals: list[VisualElement] | None = None,
        *,
        fail_visual_search: bool = False,
    ) -> None:
        self._results = results or []
        self._visuals = visuals or []
        self._fail_visual_search = fail_visual_search

    def search(self, query, *, user_id, course_id=None, material_id=None, top_k=5):
        return self._results

    def visual_search(self, query, *, user_id, course_id=None, material_id=None, top_k=5):
        if self._fail_visual_search:
            raise RuntimeError("visual backend exploded")
        return self._visuals[:top_k]


# ---------------------------------------------------------------------------
# 5B/5C/5D: extractor-level visual detection
# ---------------------------------------------------------------------------


class TestExtractorVisuals:
    def test_pdf_detects_figure_with_metadata(
        self, figure_pdf_path: Path
    ) -> None:
        from rag.extraction import extract_document

        document = extract_document(figure_pdf_path)
        assert document.visuals
        visual = document.visuals[0]
        assert visual.page == 2
        assert visual.source_location == "page 2"
        assert visual.visual_type == VISUAL_TYPE_FIGURE
        assert visual.caption == "Figure 1: Deadlock Detection Algorithm"
        assert "Deadlock Detection Algorithm" in visual.nearby_text
        assert visual.description == ""
        assert visual.pixel_analyzed is False
        assert visual.material_id is None

    def test_pptx_detects_picture_and_chart(
        self, picture_pptx_path: Path
    ) -> None:
        from rag.extraction import extract_document

        document = extract_document(picture_pptx_path)
        assert len(document.visuals) == 2
        picture, chart = document.visuals
        assert picture.slide == 1
        assert picture.source_location == "slide 1"
        assert picture.caption == "Deadlock Detection"
        assert "Deadlock Detection" in picture.nearby_text
        assert picture.visual_type == VISUAL_TYPE_IMAGE
        assert chart.slide == 2
        assert chart.visual_type == VISUAL_TYPE_CHART
        assert chart.caption == "Banker's Algorithm"

    def test_docx_detects_inline_image_with_caption(
        self, inline_image_docx_path: Path
    ) -> None:
        from rag.extraction import extract_document

        document = extract_document(inline_image_docx_path)
        assert len(document.visuals) == 1
        visual = document.visuals[0]
        assert visual.page is None
        assert visual.slide is None
        assert visual.source_location is None
        assert visual.visual_type == VISUAL_TYPE_FIGURE
        assert visual.caption == "Figure 2: Resource Allocation Graph explained."

    def test_pdf_without_images_has_no_visuals(self, pdf_path: Path) -> None:
        from rag.extraction import extract_document

        document = extract_document(pdf_path)
        assert document.visuals == ()

    def test_txt_never_has_visuals(self, txt_path: Path) -> None:
        from rag.extraction import extract_document

        document = extract_document(txt_path)
        assert document.visuals == ()


# ---------------------------------------------------------------------------
# 5E: pipeline integration
# ---------------------------------------------------------------------------


class TestPipelineIntegration:
    def test_process_document_materializes_visual_sources(
        self, figure_pdf_path: Path
    ) -> None:
        processed = process_document(
            figure_pdf_path,
            source_ref=SourceRef(
                material_id="mat-9",
                course_id="course-2",
                original_filename="deadlock_figure.pdf",
                material_title="Deadlock Notes",
                source_type=".pdf",
            ),
        )
        assert processed.visuals
        visual = processed.visuals[0]
        assert visual.material_id == "mat-9"
        assert visual.course_id == "course-2"
        assert visual.original_filename == "deadlock_figure.pdf"
        assert visual.material_title == "Deadlock Notes"
        assert visual.visual_id == "mat-9:p2"
        assert visual.page == 2

    def test_process_document_text_only_has_no_visuals(
        self, pdf_path: Path
    ) -> None:
        processed = process_document(
            pdf_path,
            source_ref=SourceRef(material_id="m1", course_id="c1"),
        )
        assert processed.visuals == ()


# ---------------------------------------------------------------------------
# 5G/5L: visual storage, search, scoping
# ---------------------------------------------------------------------------


class TestVisualStore:
    def _visuals(self) -> list[VisualElement]:
        diagram = _visual(
            "alice:m1:p2",
            material_id="m1",
            course_id="c1",
            page=2,
            source_location="page 2",
            visual_type="diagram",
            caption="Deadlock Detection Algorithm",
            nearby_text="Wait-for graph showing process cycles.",
        )
        figure_other_course = _visual(
            "alice:m2:p5",
            material_id="m2",
            course_id="c9",
            page=5,
            source_location="page 5",
            visual_type="figure",
            caption="Banker's Algorithm flow",
            nearby_text="Safety check ordered by resource need.",
        )
        return [diagram, figure_other_course]

    def test_add_and_count_visuals(
        self, embedder: DeterministicEmbedder, tmp_path: Path
    ) -> None:
        store = SqliteVectorStore(str(tmp_path / "v.db"))
        try:
            assert store.add_visuals(self._visuals(), uploaded_by="alice") == 2
            assert store.add_visuals([], uploaded_by="alice") == 0
            assert store.count_visuals() == 2
            assert (
                store.count_visuals(
                    SearchFilter(user_id="alice", course_id="c1")
                )
                == 1
            )
            assert store.count_visuals(SearchFilter(user_id="bob")) == 0
        finally:
            store.close()

    def test_search_visuals_ranks_caption_over_nearby(
        self, embedder: DeterministicEmbedder, tmp_path: Path
    ) -> None:
        store = SqliteVectorStore(str(tmp_path / "v.db"))
        try:
            store.add_visuals(self._visuals(), uploaded_by="alice")
            results = store.search_visuals(
                "deadlock", SearchFilter(user_id="alice"), top_k=5
            )
            assert [r.visual_id for r in results] == ["alice:m1:p2"]

            results = store.search_visuals(
                "banker's safe flow",
                SearchFilter(user_id="alice"),
                top_k=5,
            )
            assert [r.visual_id for r in results] == ["alice:m2:p5"]
        finally:
            store.close()

    def test_search_visuals_scoped_by_user_and_material(
        self, embedder: DeterministicEmbedder, tmp_path: Path
    ) -> None:
        store = SqliteVectorStore(str(tmp_path / "v.db"))
        try:
            store.add_visuals(self._visuals(), uploaded_by="alice")
            store.add_visuals(
                [
                    _visual(
                        "bob:m3:p1",
                        material_id="m3",
                        course_id="c1",
                        page=1,
                        source_location="page 1",
                        caption="Deadlock Detection Algorithm",
                    )
                ],
                uploaded_by="bob",
            )

            alice_hits = store.search_visuals(
                "deadlock", SearchFilter(user_id="alice"), top_k=5
            )
            assert [v.visual_id for v in alice_hits] == ["alice:m1:p2"]

            bob_hits = store.search_visuals(
                "deadlock", SearchFilter(user_id="bob"), top_k=5
            )
            assert [v.visual_id for v in bob_hits] == ["bob:m3:p1"]
            assert bob_hits[0].material_id == "m3"

            assert store.search_visuals(
                "deadlock",
                SearchFilter(user_id="bob", material_id="other"),
                top_k=5,
            ) == []
        finally:
            store.close()

    def test_search_visuals_blank_or_empty_returns_empty(
        self, embedder: DeterministicEmbedder, tmp_path: Path
    ) -> None:
        store = SqliteVectorStore(str(tmp_path / "v.db"))
        try:
            store.add_visuals(self._visuals(), uploaded_by="alice")
            assert (
                store.search_visuals("", SearchFilter(user_id="alice")) == []
            )
            assert (
                store.search_visuals("a b c", SearchFilter(user_id="alice"))
                == []
            )
        finally:
            store.close()

    def test_delete_material_removes_visuals(
        self, embedder: DeterministicEmbedder, tmp_path: Path
    ) -> None:
        store = SqliteVectorStore(str(tmp_path / "v.db"))
        try:
            store.add_visuals(self._visuals(), uploaded_by="alice")
            assert store.count_visuals() == 2
            assert store.delete_material("m2") == 0
            assert store.count_visuals() == 1
            assert (
                store.search_visuals(
                    "banker", SearchFilter(user_id="alice")
                )
                == []
            )
        finally:
            store.close()


class TestKnowledgeBaseVisuals:
    def test_index_material_persists_visuals(
        self, kb: KnowledgeBase, figure_pdf_path: Path
    ) -> None:
        kb.index_material(
            figure_pdf_path,
            source_ref=SourceRef(
                material_id="m-vis",
                course_id="c1",
                original_filename="deadlock_figure.pdf",
                material_title="Deadlock Notes",
            ),
            uploaded_by="alice",
        )
        assert kb.count_visuals() == 1
        assert kb.count_visuals(user_id="bob") == 0

    def test_visual_search_is_scoped_by_ownership(
        self, kb: KnowledgeBase, figure_pdf_path: Path
    ) -> None:
        kb.index_material(
            figure_pdf_path,
            source_ref=SourceRef(material_id="m-vis", course_id="c1"),
            uploaded_by="alice",
        )
        results = kb.visual_search(
            "deadlock detection algorithm",
            user_id="alice",
        )
        assert len(results) == 1
        assert results[0].material_id == "m-vis"
        assert kb.visual_search("deadlock", user_id="bob") == []

    def test_visual_search_course_and_material_filters(
        self, kb: KnowledgeBase, figure_pdf_path: Path
    ) -> None:
        kb.index_material(
            figure_pdf_path,
            source_ref=SourceRef(material_id="m-vis", course_id="c1"),
            uploaded_by="alice",
        )
        assert len(
            kb.visual_search(
                "deadlock", user_id="alice", course_id="c9"
            )
        ) == 0
        assert len(
            kb.visual_search(
                "deadlock", user_id="alice", material_id="other"
            )
        ) == 0

    def test_delete_material_removes_visuals(
        self, kb: KnowledgeBase, figure_pdf_path: Path
    ) -> None:
        kb.index_material(
            figure_pdf_path,
            source_ref=SourceRef(material_id="m-vis", course_id="c1"),
            uploaded_by="alice",
        )
        assert kb.count_visuals() == 1
        kb.delete_material("m-vis")
        assert kb.count_visuals() == 0


# ---------------------------------------------------------------------------
# 5F: visual-query understanding
# ---------------------------------------------------------------------------


class TestVisualIntent:
    @pytest.mark.parametrize(
        "query",
        [
            "explain the deadlock diagram",
            "show me the diagram",
            "which figure is related to deadlock",
            "explain this architecture",
            "explain the flowchart",
            "what does the chart show",
            "which graph illustrates the wait-for cycle",
            "give me a 10-mark answer and tell me which diagram from my "
            "notes should be included",
            "what diagram should I draw for deadlock",
        ],
    )
    def test_visual_queries_are_flagged(self, query: str) -> None:
        assert analyze_query(query).is_visual is True

    @pytest.mark.parametrize(
        "query",
        [
            "what is deadlock?",
            "explain deadlock prevention and avoidance",
            "write a summary of module 3",
            "hello",
            "compare TCP and UDP",
        ],
    )
    def test_text_queries_are_not_flagged(self, query: str) -> None:
        assert analyze_query(query).is_visual is False

    def test_empty_query_has_no_visual_flag(self) -> None:
        assert analyze_query("").is_visual is False


# ---------------------------------------------------------------------------
# 5G: evidence-level visual collection
# ---------------------------------------------------------------------------


class TestVisualEvidence:
    def test_collect_visual_evidence_uses_visual_search(
        self, kb: KnowledgeBase, figure_pdf_path: Path
    ) -> None:
        kb.index_material(
            figure_pdf_path,
            source_ref=SourceRef(material_id="m-vis", course_id="c1"),
            uploaded_by="alice",
        )
        visuals = collect_visual_evidence(
            kb,
            query="deadlock detection algorithm",
            user_id="alice",
            max_visuals=2,
        )
        assert len(visuals) == 1
        assert visuals[0].page == 2

    def test_collect_visual_evidence_limits(self) -> None:
        kb = _FakeKnowledgeBase(
            visuals=[_visual(visual_id=str(i)) for i in range(4)]
        )
        assert (
            len(
                collect_visual_evidence(
                    kb,
                    query="diagram",
                    user_id="alice",
                    max_visuals=2,
                )
            )
            == 2
        )

    def test_collect_visual_evidence_is_silent_on_failure(self) -> None:
        kb = _FakeKnowledgeBase(fail_visual_search=True)
        assert (
            collect_visual_evidence(
                kb, query="diagram", user_id="alice", max_visuals=2
            )
            == []
        )

    def test_collect_visual_evidence_absent_search_support(self) -> None:
        class NoVisualSearch:
            def search(self, *args, **kwargs):
                return []

        assert collect_visual_evidence(
            NoVisualSearch(), query="diagram", user_id="alice"
        ) == []


# ---------------------------------------------------------------------------
# 5H/5I/5J: grounded visual prompts + attribution
# ---------------------------------------------------------------------------


class TestVisualPrompts:
    def _result(self, text: str = "Wait-for graph scans for cycles.") -> SearchResult:
        return SearchResult(
            chunk_id="m1:0",
            material_id="m1",
            course_id="c1",
            uploaded_by="alice",
            text=text,
            score=0.9,
            metadata=ChunkMetadata(
                material_id="m1",
                course_id="c1",
                original_filename="notes.pdf",
                material_title="Deadlock Notes",
                source_type=".pdf",
                chunk_index=0,
                total_chunks=1,
                source_location="Page 1",
            ),
        )

    def test_visual_block_renders_metadata_only(self) -> None:
        system_prompt, user_prompt = build_grounded_prompt(
            "which diagram shows deadlock?",
            [self._result()],
            visuals=[_visual(source_location="page 12")],
        )
        assert "VISUAL SOURCE [Source 2]:" in user_prompt
        assert "Material: Unknown material" in user_prompt
        assert "page 12" in user_prompt
        assert "Type: diagram" in user_prompt
        assert "Caption: Deadlock Detection Algorithm" in user_prompt
        assert "Nearby text:" in user_prompt
        assert "Available source numbers for citations: [Source 1] [Source 2]." in user_prompt
        assert VISUAL_GROUNDING_NOTICE in system_prompt

    def test_visual_only_prompt_is_not_no_context(self) -> None:
        system_prompt, user_prompt = build_grounded_prompt(
            "which diagram is relevant?",
            [],
            visuals=[_visual()],
        )
        assert system_prompt != NO_CONTEXT_SYSTEM_PROMPT
        assert system_prompt.startswith(SYSTEM_PROMPT)
        assert "VISUAL SOURCE [Source 1]:" in user_prompt
        assert "Available source numbers for citations: [Source 1]." in user_prompt

    def test_no_text_and_no_visuals_keeps_no_context(self) -> None:
        system_prompt, user_prompt = build_grounded_prompt("q", [])
        assert system_prompt == NO_CONTEXT_SYSTEM_PROMPT
        assert user_prompt == "q"

    def test_visual_source_numbering_continues_after_text(self) -> None:
        from rag.context_builder import build_visual_context_block

        block = build_visual_context_block(
            [_visual(), _visual(visual_id="m1:p3", page=3)],
            start_index=3,
        )
        assert "VISUAL SOURCE [Source 3]:" in block
        assert "VISUAL SOURCE [Source 4]:" in block

    def test_empty_visual_block_is_empty(self) -> None:
        from rag.context_builder import build_visual_context_block

        assert build_visual_context_block([]) == ""

    def test_sources_from_visuals_number_after_text(self) -> None:
        text_sources = 2
        visual_sources = sources_from_visuals(
            [_visual(), _visual(visual_id="m1:p3", page=3)],
            start_index=text_sources + 1,
        )
        assert [s.source_index for s in visual_sources] == [3, 4]
        assert visual_sources[0].source_location == "page 12"
        assert visual_sources[0].material_id == "m1"


# ---------------------------------------------------------------------------
# 5I/5M: end-to-end answer behaviour
# ---------------------------------------------------------------------------


class TestVisualAnswers:
    def test_visual_question_retrieves_text_and_visual_sources(
        self,
        kb: KnowledgeBase,
        figure_pdf_path: Path,
        tmp_path: Path,
    ) -> None:
        kb.index_material(
            figure_pdf_path,
            source_ref=SourceRef(
                material_id="m-fig",
                course_id="c1",
                original_filename="deadlock_figure.pdf",
                material_title="Deadlock Notes",
            ),
            uploaded_by="alice",
            chunk_size=1000,
            overlap=0,
        )
        llm = MockProvider()
        result = answer_question(
            kb,
            llm,
            query="which diagram shows the deadlock detection algorithm",
            user_id="alice",
        )
        assert result.has_context is True
        assert "VISUAL SOURCE [Source" in llm.last_prompt
        assert VISUAL_GROUNDING_NOTICE in (llm.last_system_prompt or "")
        visual_sources = [
            s for s in result.sources if s.source_location == "page 2"
        ]
        assert visual_sources
        assert visual_sources[0].material_id == "m-fig"

    def test_normal_text_question_never_retrieves_visuals(
        self, kb: KnowledgeBase, figure_pdf_path: Path
    ) -> None:
        kb.index_material(
            figure_pdf_path,
            source_ref=SourceRef(
                material_id="m-fig",
                course_id="c1",
                original_filename="deadlock_figure.pdf",
                material_title="Deadlock Notes",
            ),
            uploaded_by="alice",
            chunk_size=1000,
            overlap=0,
        )
        llm = MockProvider()
        result = answer_question(
            kb,
            llm,
            query="what is deadlock detection",
            user_id="alice",
        )
        assert result.has_context is True
        assert "VISUAL SOURCE" not in llm.last_prompt
        assert VISUAL_GROUNDING_NOTICE not in (llm.last_system_prompt or "")

    def test_visual_question_with_no_material_stays_no_context(self) -> None:
        kb = _FakeKnowledgeBase(results=[], visuals=[])
        result = answer_question(
            kb,
            MockProvider(),
            query="which diagram shows deadlock",
            user_id="alice",
        )
        assert result.has_context is False
        assert result.sources == []
        assert "enough information" in result.answer

    def test_visual_only_evidence_still_answers(self) -> None:
        kb = _FakeKnowledgeBase(
            results=[],
            visuals=[_visual(source_location="page 12")],
        )
        result = answer_question(
            kb,
            MockProvider(),
            query="which diagram shows deadlock",
            user_id="alice",
        )
        assert result.has_context is True
        assert len(result.sources) == 1
        assert result.sources[0].source_index == 1
        assert result.sources[0].source_location == "page 12"

    def test_knowledge_base_without_visual_search_still_answers(self) -> None:
        kb = _FakeKnowledgeBase(results=[])
        result = answer_question(
            kb,
            MockProvider(),
            query="explain the deadlock diagram",
            user_id="alice",
        )
        assert result.has_context is False
        assert "enough information" in result.answer

    def test_visual_source_counts_toward_no_context_only_with_text(self) -> None:
        kb = _FakeKnowledgeBase(
            results=[
                SearchResult(
                    chunk_id="m1:0",
                    material_id="m1",
                    course_id="c1",
                    uploaded_by="alice",
                    text="Deadlock detection scans for cycles.",
                    score=0.9,
                    metadata=ChunkMetadata(
                        material_id="m1",
                        course_id="c1",
                        source_type=".pdf",
                        source_location="Page 1",
                    ),
                )
            ],
            visuals=[_visual()],
        )
        result = answer_question(
            kb,
            MockProvider(),
            query="which diagram explains the deadlock detection cycle",
            user_id="alice",
        )
        assert result.has_context is True
        assert [s.source_index for s in result.sources] == [1, 2]