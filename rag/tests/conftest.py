"""Shared fixtures that build tiny, real documents for extractor tests.

Fixture files are generated at runtime into pytest's ``tmp_path`` so no binary
assets need to be committed. Only ``reportlab`` (test-only) is used to create
real PDFs, because pypdf cannot author text PDFs.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from rag.tests.sample_data import (
    DOCX_PARAGRAPH,
    DOCX_TABLE_ROWS,
    PAGE_ONE_BODY,
    PAGE_ONE_HEADING,
    PAGE_TWO_BODY,
    PAGE_TWO_HEADING,
    SLIDE_ONE_LINES,
    SLIDE_TWO_LINES,
    SLIDE_TWO_NOTES,
    TXT_CONTENT,
)


@pytest.fixture()
def pdf_path(tmp_path: Path) -> Path:
    path = tmp_path / "lecture.pdf"
    pdf = canvas.Canvas(str(path), pagesize=letter)
    for lines in (
        [PAGE_ONE_HEADING, PAGE_ONE_BODY],
        [PAGE_TWO_HEADING, PAGE_TWO_BODY],
    ):
        y = 720
        for line in lines:
            pdf.drawString(72, y, line)
            y -= 24
        pdf.showPage()
    pdf.save()
    return path


@pytest.fixture()
def docx_path(tmp_path: Path) -> Path:
    path = tmp_path / "notes.docx"
    document = Document()
    document.add_paragraph(DOCX_PARAGRAPH)
    document.add_paragraph("Third normal form depends on every candidate key.")
    table = document.add_table(rows=len(DOCX_TABLE_ROWS), cols=2)
    for row_index, row in enumerate(DOCX_TABLE_ROWS):
        for col_index, value in enumerate(row):
            table.cell(row_index, col_index).text = value
    document.save(str(path))
    return path


@pytest.fixture()
def pptx_path(tmp_path: Path) -> Path:
    path = tmp_path / "slides.pptx"
    presentation = Presentation()
    blank_layout = presentation.slide_layouts[6]

    for lines in (SLIDE_ONE_LINES, SLIDE_TWO_LINES):
        slide = presentation.slides.add_slide(blank_layout)
        box = slide.shapes.add_textbox(
            Inches(1), Inches(1), Inches(8), Inches(4)
        )
        text_frame = box.text_frame
        for index, line in enumerate(lines):
            if index == 0:
                text_frame.text = line
            else:
                text_frame.add_paragraph().text = line

    presentation.slides[1].notes_slide.notes_text_frame.text = SLIDE_TWO_NOTES
    presentation.save(str(path))
    return path


@pytest.fixture()
def txt_path(tmp_path: Path) -> Path:
    path = tmp_path / "reading.txt"
    path.write_text(TXT_CONTENT, encoding="utf-8")
    return path


@pytest.fixture()
def empty_txt_path(tmp_path: Path) -> Path:
    path = tmp_path / "empty.txt"
    path.write_text("   \n\n  \t \n", encoding="utf-8")
    return path


@pytest.fixture()
def empty_pdf_path(tmp_path: Path) -> Path:
    path = tmp_path / "blank.pdf"
    pdf = canvas.Canvas(str(path), pagesize=letter)
    pdf.showPage()
    pdf.save()
    return path


@pytest.fixture()
def corrupt_pdf_path(tmp_path: Path) -> Path:
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"this is definitely not a pdf file %PDF nope")
    return path


@pytest.fixture()
def corrupt_docx_path(tmp_path: Path) -> Path:
    path = tmp_path / "broken.docx"
    path.write_bytes(b"PK\x03\x04 not actually a zip package")
    return path


@pytest.fixture()
def corrupt_pptx_path(tmp_path: Path) -> Path:
    path = tmp_path / "broken.pptx"
    path.write_bytes(b"PK\x03\x04 not actually a zip package")
    return path


@pytest.fixture()
def unsupported_path(tmp_path: Path) -> Path:
    path = tmp_path / "legacy.doc"
    path.write_bytes(b"\xd0\xcf\x11\xe0 legacy binary word document")
    return path
